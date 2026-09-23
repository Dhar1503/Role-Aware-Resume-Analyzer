"""
Load PDF / DOCX / TXT resumes into a layout-aware ``Document``.

Two views of the text are kept:

* ``lines``      - best-effort reading order with layout attributes (font size,
                   bold, zone). Used for section detection and extraction.
* ``ats_text``   - what a simple ATS parser sees: PDFs read strictly
                   left-to-right, top-to-bottom (so side-by-side columns get
                   interleaved); DOCX body and tables only (text boxes and
                   page headers/footers are skipped by many parsers).

Layout evidence (column gutters, tables, images, text boxes, unmapped glyphs)
is recorded for the ATS parseability checks.
"""

from __future__ import annotations

import io
import re
import statistics
import zipfile
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from pathlib import Path
from typing import List, Optional, Union

SUPPORTED_TYPES = ("pdf", "docx", "txt")
_BOLD_FONT = re.compile(r"bold|black|heavy|semibold|demi", re.IGNORECASE)
_GARBLED = re.compile(r"\(cid:\d+\)|[\ue000-\uf8ff\ufffd]")   # unmapped glyphs, icon-font PUA chars, U+FFFD


class DocumentError(ValueError):
    """Base class for documents that cannot be analysed."""


class UnsupportedFormatError(DocumentError):
    pass


class UnreadableDocumentError(DocumentError):
    pass


@dataclass
class Line:
    text: str
    page: int = 1
    zone: str = "body"              # body | table | header | footer | textbox
    size: Optional[float] = None    # font size in pt (PDF/DOCX when known)
    bold: bool = False
    style: str = ""                 # DOCX paragraph style name


@dataclass
class PageLayout:
    number: int
    width: float
    height: float
    char_count: int
    image_count: int
    image_area_ratio: float         # share of the page covered by images
    table_count: int
    column_gutters: List[float] = field(default_factory=list)   # x positions as a fraction of width


@dataclass
class Document:
    filename: str
    file_type: str
    size_bytes: int
    lines: List[Line]
    ats_text: str
    pages: List[PageLayout] = field(default_factory=list)
    page_count: Optional[int] = None
    page_count_estimated: bool = False
    table_count: int = 0
    image_count: int = 0
    textbox_words: int = 0
    header_footer_text: str = ""
    docx_columns: int = 1
    garbled_chars: int = 0          # unmapped glyphs inside text (content may be lost)
    garbled_symbols: int = 0        # standalone unmapped glyphs: bullets, icons (cosmetic)
    reading_order_similarity: Optional[float] = None   # PDF: agreement of two independent extractors

    @property
    def text(self) -> str:
        """All text a human reader would see, in reading order."""
        return "\n".join(line.text for line in self.lines)

    @property
    def ats_lines(self) -> List[Line]:
        """Lines an ATS is expected to read (body and tables)."""
        return [line for line in self.lines if line.zone in ("body", "table")]

    @property
    def word_count(self) -> int:
        return len(self.text.split())


# --------------------------------------------------------------------------
# Entry point
# --------------------------------------------------------------------------

def load_document(source: Union[str, Path, bytes], filename: Optional[str] = None) -> Document:
    """Load a resume from a path or raw bytes. ``filename`` is required with bytes."""
    if isinstance(source, (str, Path)):
        path = Path(source)
        filename = filename or path.name
        try:
            data = path.read_bytes()
        except OSError as exc:
            raise UnreadableDocumentError(f"Could not read '{filename}': {exc.strerror}") from None
    else:
        if not filename:
            raise ValueError("filename is required when loading from bytes")
        data = source

    ext = Path(filename).suffix.lower().lstrip(".")
    if ext == "doc":
        raise UnsupportedFormatError("Legacy .doc files cannot be read reliably. Save the resume as .docx or PDF.")
    if ext in ("png", "jpg", "jpeg", "gif", "bmp", "tiff", "webp"):
        raise UnsupportedFormatError("Image files have no text layer an ATS can read. Upload a PDF or DOCX.")
    if ext not in SUPPORTED_TYPES:
        raise UnsupportedFormatError(f"Unsupported file type '.{ext or '?'}'. Upload a PDF, DOCX or TXT file.")
    if not data:
        raise UnreadableDocumentError(f"'{filename}' is empty.")

    loader = {"pdf": _load_pdf, "docx": _load_docx, "txt": _load_txt}[ext]
    return loader(data, filename)


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

def _load_pdf(data: bytes, filename: str) -> Document:
    if not data.lstrip()[:5].startswith(b"%PDF"):
        raise UnreadableDocumentError(f"'{filename}' has a .pdf extension but is not a PDF file.")
    import pdfplumber
    import pymupdf

    try:
        mu = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:  # pymupdf raises several internal types for corrupt files
        raise UnreadableDocumentError(f"'{filename}' could not be opened as a PDF ({exc}).") from None
    if mu.needs_pass:
        raise UnreadableDocumentError(f"'{filename}' is password-protected. Remove the password and upload again.")

    lines: List[Line] = []
    for page_no, page in enumerate(mu, start=1):
        blocks: List[List[tuple]] = []
        for block in page.get_text("dict")["blocks"]:
            raw: List[tuple] = []       # (Line, bbox)
            for ln in block.get("lines", []):
                spans = [s for s in ln["spans"] if s["text"].strip()]
                if not spans:
                    continue
                text = " ".join("".join(s["text"] for s in ln["spans"]).split())
                size = max(s["size"] for s in spans)
                bold = all((s["flags"] & 16) or _BOLD_FONT.search(s["font"]) for s in spans)
                raw.append((Line(text, page_no, "body", round(size, 1), bool(bold)), ln["bbox"]))
            if raw:
                blocks.append(raw)
        leading = _page_leading(blocks)
        for raw in blocks:
            lines.extend(_join_wrapped(raw, leading))
    mu.close()

    pages: List[PageLayout] = []
    naive: List[str] = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page_no, page in enumerate(pdf.pages, start=1):
                page_text = page.extract_text() or ""
                naive.append(page_text)
                area = float(page.width * page.height) or 1.0
                img_area = sum(max(0.0, (im["x1"] - im["x0"]) * (im["bottom"] - im["top"])) for im in page.images)
                tables = [t for t in page.find_tables() if len(t.rows) >= 2 and len(t.rows[0].cells) >= 2]
                # Words inside tables are reported by the table check, not mistaken for columns.
                words = [w for w in page.extract_words(use_text_flow=False, keep_blank_chars=False)
                         if not any(t.bbox[0] <= w["x0"] and w["x1"] <= t.bbox[2] and t.bbox[1] <= w["top"]
                                    and w["bottom"] <= t.bbox[3] for t in tables)]
                pages.append(PageLayout(
                    number=page_no, width=float(page.width), height=float(page.height),
                    char_count=len(page_text.replace(" ", "").replace("\n", "")),
                    image_count=len(page.images), image_area_ratio=round(min(1.0, img_area / area), 3),
                    table_count=len(tables), column_gutters=detect_column_gutters(words, float(page.width))))
    except Exception as exc:
        raise UnreadableDocumentError(f"'{filename}' could not be parsed as a PDF ({exc}).") from None

    ats_text = "\n".join(naive)
    content, bullets = count_garbled(ats_text)
    return Document(
        filename=filename, file_type="pdf", size_bytes=len(data), lines=lines, ats_text=ats_text,
        pages=pages, page_count=len(pages),
        table_count=sum(p.table_count for p in pages), image_count=sum(p.image_count for p in pages),
        garbled_chars=content, garbled_symbols=bullets,
        reading_order_similarity=word_order_similarity(" ".join(l.text for l in lines), ats_text),
    )


def detect_column_gutters(words: List[dict], page_width: float) -> List[float]:
    """Find vertical strips no word crosses, with substantial text on both sides.

    A single-column resume always has full-width lines (bullets, summaries)
    crossing any candidate strip, even when dates are right-aligned. In a real
    multi-column layout one strip stays clear for most of the page height.
    Returns gutter x positions as fractions of page width.
    """
    words = [w for w in words if w["text"].strip()]
    if len(words) < 30:
        return []
    top = min(w["top"] for w in words)
    bottom = max(w["bottom"] for w in words)
    height = bottom - top
    if height <= 0:
        return []

    best: Optional[tuple] = None   # (clear_fraction, x)
    x = page_width * 0.2
    while x <= page_width * 0.75:
        crossing = sorted((w["top"], w["bottom"]) for w in words if w["x0"] < x < w["x1"])
        # Longest vertical run with no crossing word.
        longest, cursor, run_start = 0.0, top, top
        for w_top, w_bottom in crossing:
            if w_top > cursor:
                if w_top - cursor > longest:
                    longest, run_start = w_top - cursor, cursor
            cursor = max(cursor, w_bottom)
        if bottom - cursor > longest:
            longest, run_start = bottom - cursor, cursor
        clear = longest / height
        if clear >= 0.5:
            run_end = run_start + longest
            inside = [w for w in words if w["top"] >= run_start and w["bottom"] <= run_end]
            left = [w for w in inside if w["x1"] <= x]
            right = [w for w in inside if w["x0"] >= x]
            # Real columns have text running down both sides of the gap, not
            # just a few right-aligned dates beside an empty area.
            if (min(len(left), len(right)) >= 10
                    and _vertical_coverage(left) >= 0.3 * longest
                    and _vertical_coverage(right) >= 0.3 * longest):
                if best is None or clear > best[0]:
                    best = (clear, x)
        x += 2.0
    return [round(best[1] / page_width, 3)] if best else []


_WRAP_TOLERANCE = 12.0      # pt: how close to the block's right edge counts as "full line"
_HAS_DATE_HINT = re.compile(r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s*['’]?\s*(?:19|20)?\d{2}|\b(?:19|20)\d{2}\s*(?:-|–|—|to)\s*(?:(?:19|20)\d{2}|present|current)", re.IGNORECASE)
_STARTS_ITEM = re.compile(r"^\s*(?:[-*•‣▪●◦⁃·]|\(cid:\d+\)|"
                          r"[-]|\d{1,2}[.)]\s)")


def _page_leading(blocks: List[List[tuple]]) -> float:
    """The page's tightest line spacing, i.e. the gap inside a wrapped paragraph."""
    gaps = sorted(round(cur[1] - prev[3], 1)
                  for raw in blocks for (_, prev), (_, cur) in zip(raw, raw[1:])
                  if -2 <= cur[1] - prev[3] <= 40)
    if not gaps:
        return 0.0
    return gaps[max(0, int(len(gaps) * 0.1) - 1)]      # 10th percentile


def _join_wrapped(raw: List[tuple], leading: float = 0.0) -> List[Line]:
    """Rejoin lines that are only line-wrapped continuations of the one above.

    A continuation follows a line that filled the block's width, breaks
    mid-sentence, is indented at least as far as its parent, and never starts a
    new item or carries a date of its own. Without this, one long bullet looks
    like two entries and every count built on entries (projects, publications,
    certifications) is wrong; with a rule any looser, the next entry's header
    gets swallowed by the bullet above it.
    """
    if not raw:
        return []
    right_edge = max(bbox[2] for _, bbox in raw)
    out: List[Line] = []
    previous: Optional[tuple] = None        # (x0, filled_to_edge, bottom)
    for line, (x0, top, x1, bottom) in raw:
        tight = previous is not None and (top - previous[2]) <= leading * 1.4 + 0.6
        merge = (out and previous and previous[1] and tight
                 and x0 >= previous[0] - 1
                 and not _STARTS_ITEM.match(line.text)
                 and not _HAS_DATE_HINT.search(line.text)
                 and out[-1].size == line.size and out[-1].bold == line.bold
                 and not out[-1].text.rstrip().endswith((".", ":", ";", "!", "?")))
        if merge:
            out[-1].text = f"{out[-1].text} {line.text}".strip()
        else:
            out.append(line)
        previous = (x0, x1 >= right_edge - _WRAP_TOLERANCE, bottom)
    return out


def _vertical_coverage(words: List[dict]) -> float:
    """Total height covered by the union of the words' vertical extents."""
    total, end = 0.0, float("-inf")
    for top, bottom in sorted((w["top"], w["bottom"]) for w in words):
        if top > end:
            total += bottom - top
            end = bottom
        elif bottom > end:
            total += bottom - end
            end = bottom
    return total


def word_order_similarity(a: str, b: str) -> Optional[float]:
    """How closely two extractions agree on word order (1.0 = identical).

    Compared at word level because extractors legitimately differ on where
    lines break (e.g. a right-aligned date on its own line or not). Side-by-side
    columns interleave words, which drives the ratio down.
    """
    tokenize = lambda s: [w for w in re.findall(r"[a-z0-9]+", s.lower()) if len(w) > 1]
    wa, wb = tokenize(a), tokenize(b)
    if len(wa) < 20 or len(wb) < 20:
        return None
    return round(SequenceMatcher(None, wa, wb, autojunk=False).ratio(), 3)


def count_garbled(text: str) -> tuple:
    """(content, symbol) counts of unmapped glyphs.

    A glyph standing alone as a token is a bullet or icon (cosmetic); one
    attached to letters, as in "de(cid:12)ned", means characters of real text
    were lost.
    """
    content = symbols = 0
    for m in _GARBLED.finditer(text):
        before = text[m.start() - 1] if m.start() else " "
        after = text[m.end()] if m.end() < len(text) else " "
        if before.isspace() and after.isspace():
            symbols += 1
        else:
            content += 1
    return content, symbols


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _load_docx(data: bytes, filename: str) -> Document:
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise UnreadableDocumentError(f"'{filename}' has a .docx extension but is not a Word document.")
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        doc = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise UnreadableDocumentError(f"'{filename}' could not be opened as a Word document ({exc}).") from None

    def para_line(p: Paragraph, zone: str) -> Optional[Line]:
        text = " ".join(p.text.split())
        if not text:
            return None
        runs = [r for r in p.runs if r.text.strip()]
        style_font = p.style.font if p.style is not None else None
        bold = bool(runs) and all(r.bold or (r.bold is None and style_font is not None and style_font.bold) for r in runs)
        sizes = [r.font.size.pt for r in runs if r.font.size is not None]
        if not sizes and style_font is not None and style_font.size is not None:
            sizes = [style_font.size.pt]
        return Line(text, 1, zone, max(sizes) if sizes else None, bool(bold), p.style.name if p.style is not None else "")

    lines: List[Line] = []
    ats_parts: List[str] = []
    tables = 0
    for child in doc.element.body.iterchildren():
        if child.tag == f"{_W}p":
            line = para_line(Paragraph(child, doc), "body")
            if line:
                lines.append(line)
                ats_parts.append(line.text)
        elif child.tag == f"{_W}tbl":
            tables += 1
            for row in Table(child, doc).rows:
                cells: List[str] = []
                for cell in row.cells:
                    for p in cell.paragraphs:
                        line = para_line(p, "table")
                        if line and (not lines or lines[-1].text != line.text):   # merged cells repeat
                            lines.append(line)
                            cells.append(line.text)
                if cells:
                    ats_parts.append(" | ".join(dict.fromkeys(cells)))

    # Text boxes: python-docx skips them; many ATS do too. Word stores each box
    # twice (DrawingML + VML fallback), so de-duplicate.
    box_texts: List[str] = []
    for box in doc.element.body.iter(f"{_W}txbxContent"):
        for p in box.iter(f"{_W}p"):
            text = " ".join("".join(t.text or "" for t in p.iter(f"{_W}t")).split())
            if text and text not in box_texts:
                box_texts.append(text)
    lines.extend(Line(t, 1, "textbox") for t in box_texts)

    hf: List[str] = []
    zones = {"header": [], "footer": []}
    for section in doc.sections:
        for zone, part in (("header", section.header), ("footer", section.footer)):
            if part.is_linked_to_previous and section is not doc.sections[0]:
                continue
            paragraphs = list(part.paragraphs) + [p for t in part.tables for row in t.rows
                                                  for cell in row.cells for p in cell.paragraphs]
            for p in paragraphs:
                line = para_line(p, zone)
                if line and line.text not in hf:
                    hf.append(line.text)
                    zones[zone].append(line)
    lines = zones["header"] + lines + zones["footer"]

    columns = 1
    for cols in doc.element.body.iter(f"{_W}cols"):
        num = cols.get(f"{_W}num")
        if num and num.isdigit():
            columns = max(columns, int(num))

    images = len(doc.inline_shapes) + len(list(doc.element.body.iter(
        "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}anchor")))

    # Word records the real page count when it saves, but files generated from
    # templates carry a stale value, so only trust it when it is plausible.
    words = sum(len(l.text.split()) for l in lines)
    estimate = max(1, round(words / 500 + 0.49))
    recorded = _docx_page_count(data)
    if recorded is not None and abs(recorded - estimate) <= 1:
        page_count, estimated = recorded, False
    else:
        page_count, estimated = estimate, True

    body_text = "\n".join(ats_parts)
    return Document(
        filename=filename, file_type="docx", size_bytes=len(data), lines=lines, ats_text=body_text,
        page_count=page_count, page_count_estimated=estimated, table_count=tables, image_count=images,
        textbox_words=sum(len(t.split()) for t in box_texts), header_footer_text="\n".join(hf),
        docx_columns=columns, garbled_chars=count_garbled("\n".join(l.text for l in lines))[0],
    )


def _docx_page_count(data: bytes) -> Optional[int]:
    """Word records the page count in docProps/app.xml when it saves; generated files may not."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            xml = z.read("docProps/app.xml").decode("utf-8", "ignore")
    except (KeyError, zipfile.BadZipFile):
        return None
    m = re.search(r"<Pages>(\d+)</Pages>", xml)
    return int(m.group(1)) if m and int(m.group(1)) > 0 else None


# --------------------------------------------------------------------------
# TXT
# --------------------------------------------------------------------------

def _load_txt(data: bytes, filename: str) -> Document:
    # UTF-16 only with a byte-order mark: without one, almost any bytes "decode" as UTF-16.
    encodings = ("utf-16",) if data[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig", "cp1252")
    for encoding in encodings:
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = data.decode("utf-8", errors="replace")
    if "\x00" in text:
        raise UnreadableDocumentError(f"'{filename}' does not look like a text file.")
    lines = [Line(" ".join(raw.split())) for raw in text.splitlines() if raw.strip()]
    if not lines:
        raise UnreadableDocumentError(f"'{filename}' contains no text.")
    words = sum(len(l.text.split()) for l in lines)
    return Document(
        filename=filename, file_type="txt", size_bytes=len(data), lines=lines,
        ats_text="\n".join(l.text for l in lines), page_count=max(1, round(words / 500 + 0.49)),
        page_count_estimated=True, garbled_chars=count_garbled(text)[0],
    )


def body_font_size(lines: List[Line]) -> Optional[float]:
    """Most common font size weighted by characters (the body text size)."""
    weighted = [l.size for l in lines if l.size for _ in range(max(1, len(l.text) // 10))]
    return statistics.mode(weighted) if weighted else None
