"""
Render a layout as an ATS-safe DOCX.

Same constraints as the PDF: body paragraphs only - no tables, text boxes,
headers/footers or images - so every word sits where a parser reads it. Word's
built-in Heading styles are used (parsers recognise them) but recoloured to
black so the file still looks like a resume rather than a template.
"""

from __future__ import annotations

import io

from .layout import Layout

FONT = "Calibri"
BODY_PT = 10.5
HEADING_PT = 12


def render_docx(layout: Layout) -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    document = docx.Document()
    for section in document.sections:
        section.top_margin = section.bottom_margin = Pt(36)
        section.left_margin = section.right_margin = Pt(40)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    normal.paragraph_format.space_after = Pt(2)

    def heading(text: str) -> None:
        paragraph = document.add_heading(text, level=1)
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(2)
        for run in paragraph.runs:
            run.font.name = FONT
            run.font.size = Pt(HEADING_PT)
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.font.bold = True

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run(layout.name)
    run.bold = True
    run.font.size = Pt(18)
    if layout.contact:
        contact = document.add_paragraph(layout.contact)
        contact.runs[0].font.size = Pt(9)

    for section in layout.sections:
        heading(section.heading)
        for block in section.blocks:
            if block.kind == "para" and block.text:
                document.add_paragraph(block.text)
            elif block.kind == "entry":
                paragraph = document.add_paragraph()
                paragraph.add_run(block.text).bold = True
                for bullet in block.bullets:
                    document.add_paragraph(bullet, style="List Bullet")
            elif block.kind == "bullets":
                for bullet in block.bullets:
                    document.add_paragraph(bullet, style="List Bullet")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
